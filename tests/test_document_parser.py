from __future__ import annotations

import io
import pytest

from docx import Document
from services.document_parser import DocumentParser, DocumentParserError


def test_unsupported_file_extension():
    with pytest.raises(DocumentParserError, match="Unsupported file format"):
        DocumentParser("malicious.exe", b"fake binary content")


def test_markdown_resume_parser():
    md_content = """# Alex Rivera
Email: alex@example.com
Phone: (555) 123-4567
GitHub: https://github.com/demo-user
LinkedIn: https://linkedin.com/in/alex-rivera

## Skills
Proficient in Python, SQL, Pandas, Power BI, Excel, and Machine Learning.

## Education
Bachelor of Science in Computer Science
Sample University

## Experience
Data Analyst Intern at Analytics Corp

## Projects
Project:
Customer Churn Predictor
Built a machine learning pipeline using Python and Scikit-learn.

## Certificates
Coursera Data Science Specialization
"""
    parser = DocumentParser("resume.md", md_content.encode("utf-8"))
    res = parser.parse()

    assert res["profile"]["full_name"] == "# Alex Rivera"
    assert res["profile"]["email"] == "alex@example.com"
    assert res["profile"]["github_username"] == "demo-user"
    
    skill_names = [s["name"] for s in res["skills"]]
    assert "Python" in skill_names
    assert "SQL" in skill_names
    assert "Pandas" in skill_names
    assert "Power BI" in skill_names

    for s in res["skills"]:
        assert s["verified"] is False

    assert len(res["education"]) > 0
    assert res["education"][0]["confidence"] == "uncertain"


def test_csv_tabular_parser():
    csv_content = """company,role,status,deadline
Acme Corp,Data Analyst Intern,applied,2026-10-01
Beta Analytics,AI Research Intern,wishlist,2026-11-15
"""
    parser = DocumentParser("applications.csv", csv_content.encode("utf-8"))
    res = parser.parse()
    assert "tabular_data" in res
    assert res["tabular_data"]["type"] == "applications"
    assert res["tabular_data"]["count"] == 2
    assert res["tabular_data"]["records"][0]["company"] == "Acme Corp"


def test_docx_parser():
    doc = Document()
    doc.add_paragraph("Alex Rivera")
    doc.add_paragraph("Email: user@example.com")
    doc.add_paragraph("Skills: Python, Streamlit, PostgreSQL")
    bio = io.BytesIO()
    doc.save(bio)

    parser = DocumentParser("test.docx", bio.getvalue())
    res = parser.parse()
    assert res["profile"]["email"] == "user@example.com"
    skill_names = [s["name"] for s in res["skills"]]
    assert "Python" in skill_names
    assert "Streamlit" in skill_names
